#!/usr/bin/env python3
"""Recompute per-rater keep rates and pairwise Cohen's kappa for every screening stage.

Every number in Methodology/Screening_Reliability_Statistics.md comes from this script,
read directly from the ground-truth decision files (no Zotero calls):

  Pass 1 pilot   slr-phase-1-2/ssrn-decisions.xlsx   (sheets: decisions, chatgpt, gemini;
                 the CSV export drops the human column -- the xlsx is mandatory)
  Pass 1 non-SSRN slr-phase-1-2/nonssrn-decisions-2026-05-25.csv
  Pass 2 Trust Check  slr-phase-1-2/phase2/verification/trust_check.csv
  Stage 3 QA     slr-tools/stage3/work/qa/{qa_master_250,codex_out_250,gemini_out_250,human_review_50}.csv

Needs openpyxl (use slr-tools/stage6/.venv/bin/python).
"""

from __future__ import annotations

import collections
import csv
import itertools
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = lambda *a: os.path.join(ROOT, *a)


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = collections.Counter(a), collections.Counter(b)
    pe = sum(ca[c] * cb[c] for c in set(a) | set(b)) / n / n
    return (po - pe) / (1 - pe), po, n


VALID = {"keep", "maybe", "discard", "core", "context"}
TYPOS = {"dsicard": "discard"}   # one human cell in ssrn-decisions.xlsx; intent unambiguous


def norm(v):
    v = str(v or "").strip().lower()
    v = TYPOS.get(v, v)
    return v if v in VALID else None


def rates(d, cats, keys=None):
    keys = [k for k in (keys if keys is not None else d) if d.get(k)]
    c = collections.Counter(d[k] for k in keys)
    n = len(keys)
    return n, {x: c[x] / n for x in cats}


def pairs(R, collapse):
    out = []
    for a, b in itertools.combinations(R, 2):
        ks = sorted(k for k in set(R[a]) & set(R[b]) if R[a][k] and R[b][k])
        if len(ks) < 5:
            continue
        x, y = [R[a][k] for k in ks], [R[b][k] for k in ks]
        row = {"pair": f"{a} / {b}", "n": len(ks)}
        row["k_full"], row["po_full"], _ = kappa(x, y)
        for name, f in collapse.items():
            row["k_" + name], row["po_" + name], _ = kappa([f(v) for v in x], [f(v) for v in y])
        out.append(row)
    return out


def show(title, R, cats, collapse, common=True):
    print(f"\n## {title}")
    for r, d in R.items():
        n, rt = rates(d, cats)
        print(f"  {r:10s} n={n:5d} " + " ".join(f"{c}={rt[c]:6.1%}" for c in cats))
    if common:
        ks = [k for k in set.intersection(*(set(d) for d in R.values())) if all(R[r].get(k) for r in R)]
        if ks:
            print(f"  -- common subset n={len(ks)}")
            for r, d in R.items():
                _, rt = rates(d, cats, ks)
                print(f"  {r:10s} " + " ".join(f"{c}={rt[c]:6.1%}" for c in cats))
    for p in pairs(R, collapse):
        extra = "  ".join(f"k[{n}]={p['k_' + n]:6.3f} (Po {p['po_' + n]:.0%})" for n in collapse)
        print(f"  {p['pair']:22s} n={p['n']:5d}  k[full]={p['k_full']:6.3f} (Po {p['po_full']:.0%})  {extra}")


P1_COLLAPSE = {"pass|discard": lambda v: "d" if v == "discard" else "p",
               "keep|not": lambda v: "k" if v == "keep" else "n"}
S3_COLLAPSE = {"keep|discard": lambda v: "d" if v == "discard" else "k",
               "core|not": lambda v: "c" if v == "core" else "n"}


def pass1_ssrn():
    import openpyxl
    wb = openpyxl.load_workbook(P("slr-phase-1-2", "ssrn-decisions.xlsx"), read_only=True)
    rows = list(wb["decisions"].iter_rows(values_only=True))
    h = list(rows[0])
    ik, hd = h.index("item_key"), h.index("human_decision")
    R = {"claude": {}, "human": {}, "chatgpt": {}, "gemini": {}}
    for r in rows[1:]:   # 3,869 rows, 6 duplicated keys -> 3,863 unique items
        r = list(r); r[ik] = str(r[ik])
        R["claude"][r[ik]] = norm(r[h.index("decision")])
        R["human"][r[ik]] = norm(r[hd])
    for s in ("chatgpt", "gemini"):
        for r in list(wb[s].iter_rows(values_only=True))[1:]:
            if r[0] and norm(r[1]):
                R[s][str(r[0])] = norm(r[1])
    return R


def pass1_nonssrn():
    R = {"claude": {}, "human": {}, "chatgpt": {}, "gemini": {}}
    for r in csv.DictReader(open(P("slr-phase-1-2", "nonssrn-decisions-2026-05-25.csv"))):
        for k in R:
            R[k][r["item_key"]] = norm(r[f"{k}_decision"])
    return R


def pass2_trust():
    R = {"sonnet/opus": {}, "human": {}}
    for r in csv.DictReader(open(P("slr-phase-1-2", "phase2", "verification", "trust_check.csv"))):
        R["sonnet/opus"][r["item_key"]] = norm(r["decision"])
        R["human"][r["item_key"]] = norm(r["human_decision"])
    return R


def stage3():
    qa = P("slr-tools", "stage3", "work", "qa")
    load = lambda f: {r["item_key"]: norm(r["bin"]) for r in csv.DictReader(open(os.path.join(qa, f)))}
    return {"opus": load("qa_master_250.csv"), "gpt-5.5": load("codex_out_250.csv"),
            "gemini": load("gemini_out_250.csv"), "human": load("human_review_50.csv")}


def main():
    cats = ["keep", "maybe", "discard"]
    ssrn, non = pass1_ssrn(), pass1_nonssrn()
    show("Pass 1 pilot (SSRN)", ssrn, cats, P1_COLLAPSE)
    show("Pass 1 non-SSRN", non, cats, P1_COLLAPSE)
    pooled = {r: {**ssrn[r], **non[r]} for r in ssrn}
    show("Pass 1 pooled", pooled, cats, P1_COLLAPSE, common=False)
    show("Pass 2 Trust Check", pass2_trust(), cats, {})
    S3 = stage3()
    show("Stage 3 QA (all 250 / human 50)", S3, ["core", "context", "discard"], S3_COLLAPSE, common=False)
    h50 = [k for k in S3["human"] if S3["human"][k]]
    print("  -- on the human-50 subset")
    for r, d in S3.items():
        n, rt = rates(d, ["core", "context", "discard"], h50)
        print(f"  {r:10s} n={n:3d} " + " ".join(f"{c}={rt[c]:6.1%}" for c in rt))


if __name__ == "__main__":
    main()
